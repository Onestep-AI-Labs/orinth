"""Phase 11: data recipes — extraction, chunking, generation, commit, settings."""

import io
import json
from pathlib import Path

import docx
import pytest
from pypdf import PdfWriter

from app.core.config import Settings
from app.core.storage import Storage
from app.schemas import (
    RecipeCommitRequest,
    RecipeCreate,
    RecipeGenerateRequest,
    RecipeGenerationSettings,
    RecipeRecordCreate,
    RecipeRecordUpdate,
)
from app.services.datasets import DatasetService
from app.services.recipes import RecipeService
from app.services.recipes import generation as gen
from app.services.recipes import openrouter
from app.services.recipes.chunking import chunk_sources
from app.services.recipes.extraction import extract_text
from app.services.recipes.service import RecipeService as ServiceCls
from app.services.settings import SettingsService


class InlineExecutor:
    """Runs submitted work synchronously so generation completes in-test."""

    def submit(self, fn, *args, **kwargs):
        fn(*args, **kwargs)
        return None


def make_service(settings: Settings, executor=None) -> tuple[RecipeService, DatasetService]:
    storage = Storage(settings)
    storage.ensure()
    datasets = DatasetService(settings, storage)
    service = RecipeService(settings, storage, datasets, executor or InlineExecutor())
    return service, datasets


def seed_recipe(service: RecipeService, text: str, output_format: str = "instruction_jsonl") -> str:
    recipe = service.create_recipe(
        RecipeCreate(name="Docs", output_format=output_format)
    )
    root = service.storage.recipes / recipe.id
    (root / "sources").mkdir(parents=True, exist_ok=True)
    source_id = "src-seed01"
    (root / "sources" / f"{source_id}.txt").write_text(text, encoding="utf-8")
    manifest = json.loads((root / "manifest.json").read_text())
    manifest["sources"].append(
        {
            "id": source_id,
            "filename": "seed.txt",
            "media_type": "txt",
            "characters": len(text),
            "pages": None,
            "excluded": not text.strip(),
            "warnings": [],
        }
    )
    (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return recipe.id


# ---- extraction -------------------------------------------------------------


def test_extract_plain_and_markdown():
    result = extract_text("notes.md", b"# Title\n\nBody text.", ".md")
    assert "Title" in result.text
    assert result.characters > 0


def test_extract_csv_maps_columns_to_lines():
    data = b"question,answer\nWhat is 2+2?,4\n"
    result = extract_text("qa.csv", data, ".csv")
    assert "question: What is 2+2?" in result.text
    assert "answer: 4" in result.text


def test_extract_jsonl_stringifies_rows():
    data = b'{"instruction": "Hi", "output": "There"}\nnot-json\n'
    result = extract_text("rows.jsonl", data, ".jsonl")
    assert "instruction: Hi" in result.text
    assert any("malformed" in w for w in result.warnings)


def test_extract_docx_reads_paragraphs():
    document = docx.Document()
    document.add_paragraph("Alpha paragraph.")
    document.add_paragraph("Beta paragraph.")
    buffer = io.BytesIO()
    document.save(buffer)
    result = extract_text("doc.docx", buffer.getvalue(), ".docx")
    assert "Alpha paragraph." in result.text
    assert "Beta paragraph." in result.text


def test_extract_scanned_pdf_warns_and_excludes():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    result = extract_text("scan.pdf", buffer.getvalue(), ".pdf")
    assert result.pages == 1
    assert result.text == ""
    assert any("no extractable text" in w.lower() for w in result.warnings)


# ---- chunking ---------------------------------------------------------------


def test_chunk_overlap_and_boundaries():
    text = "\n\n".join(f"Paragraph number {i} with some filler words." for i in range(30))
    chunks = chunk_sources([("s1", text)], chunk_size=200, chunk_overlap=50)
    assert len(chunks) > 1
    assert all(chunk.source_id == "s1" for chunk in chunks)
    assert [chunk.index for chunk in chunks] == list(range(len(chunks)))
    # Overlap: each chunk starts before the previous chunk ended.
    for previous, current in zip(chunks, chunks[1:], strict=False):
        assert current.offset < previous.offset + len(previous.text)


def test_single_short_chunk():
    chunks = chunk_sources([("s1", "short text")], chunk_size=3000, chunk_overlap=200)
    assert len(chunks) == 1
    assert chunks[0].text == "short text"


# ---- rule-based generation --------------------------------------------------


def test_rule_records_instruction_qa():
    records = gen.rule_records(
        "What is photosynthesis?\nIt converts light into energy.", "instruction_jsonl", "qa", 3
    )
    assert records
    assert records[0]["instruction"] == "What is photosynthesis?"
    assert records[0]["output"]


def test_rule_records_chat_shape():
    records = gen.rule_records(
        "Mitochondria\n\nThe mitochondria is the powerhouse of the cell.",
        "chat_jsonl",
        "conversation",
        3,
    )
    assert records
    messages = records[0]["messages"]
    roles = {m["role"] for m in messages}
    assert {"user", "assistant"} <= roles


def test_rules_generation_end_to_end(settings: Settings):
    service, _ = make_service(settings)
    text = "\n\n".join(f"Section {i}. Detail sentence about topic {i}." for i in range(5))
    recipe_id = seed_recipe(service, text)
    recipe = service.generate(recipe_id, RecipeGenerateRequest(mode="rules"))
    assert recipe.status in {"generating", "ready"}
    final = service.get_recipe(recipe_id)
    assert final.status == "ready"
    assert final.record_count > 0
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert all(record.generator == "rules" for record in page.records)


# ---- LLM path: repair + fallback (mocked OpenRouter) ------------------------


def _llm_settings(make_settings) -> Settings:
    return make_settings(OPENROUTER_API_KEY="sk-test-key", OPENROUTER_MODEL="test/model")


def test_llm_path_valid_json(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)
    payload = json.dumps([{"instruction": "Q1", "output": "A1"}])
    monkeypatch.setattr(openrouter, "chat_completion", lambda *a, **k: payload)
    recipe_id = seed_recipe(service, "Some passage of text to generate from.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert page.total >= 1
    assert page.records[0].generator == "llm"
    assert page.records[0].record["instruction"] == "Q1"


def test_llm_repair_retry_then_valid(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)
    calls = {"n": 0}

    def fake(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            return "not json at all"
        return json.dumps([{"instruction": "Repaired", "output": "Yes"}])

    monkeypatch.setattr(openrouter, "chat_completion", fake)
    recipe_id = seed_recipe(service, "Passage.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert calls["n"] == 2  # one initial + one repair
    assert page.records[0].generator == "llm"
    assert page.records[0].record["instruction"] == "Repaired"


def test_llm_double_malformed_falls_back_to_rules(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)
    monkeypatch.setattr(openrouter, "chat_completion", lambda *a, **k: "still not json")
    recipe_id = seed_recipe(service, "A meaningful passage about a subject.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    final = service.get_recipe(recipe_id)
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert page.total > 0
    assert all(record.generator == "rules" for record in page.records)
    assert any("used rules" in w for w in final.warnings)


def test_llm_auth_error_flips_run_to_rules(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)

    def boom(*args, **kwargs):
        raise openrouter.OpenRouterAuthError("rejected")

    monkeypatch.setattr(openrouter, "chat_completion", boom)
    text = "\n\n".join(f"Paragraph {i} with content." for i in range(6))
    recipe_id = seed_recipe(service, text)
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    final = service.get_recipe(recipe_id)
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert all(record.generator == "rules" for record in page.records)
    assert any("rejected" in w.lower() for w in final.warnings)


def test_llm_mode_without_key_rejected(settings: Settings):
    service, _ = make_service(settings)
    recipe_id = seed_recipe(service, "Passage.")
    with pytest.raises(Exception) as exc:
        service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    assert "OpenRouter" in str(exc.value)


# ---- cancel semantics -------------------------------------------------------


def test_cancel_between_chunks_keeps_partial(settings: Settings):
    service, _ = make_service(settings)
    text = "\n\n".join(f"Chunk paragraph {i} with filler words here." for i in range(40))
    recipe_id = seed_recipe(service, text)
    chunks = chunk_sources([("s", text)], 200, 50)
    assert len(chunks) > 2

    class CancelAfterFirst:
        def __init__(self):
            self.checks = 0

        def is_set(self):
            self.checks += 1
            return self.checks > 1  # process one chunk, then cancel

        def set(self):
            pass

    root = service.storage.recipes / recipe_id
    service._run_generation(
        recipe_id,
        [("s", text)],
        RecipeGenerationSettings(chunk_size=200, chunk_overlap=50, records_per_chunk=2),
        "instruction_jsonl",
        False,
        None,
        None,
        CancelAfterFirst(),
    )
    final = service.get_recipe(recipe_id)
    assert final.status == "ready"
    assert 0 < final.record_count  # partial records survived
    assert final.record_count < len(chunks) * 2
    assert (root / "records.jsonl").exists()


# ---- record review ----------------------------------------------------------


def test_edit_record_keeps_provenance(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)
    monkeypatch.setattr(
        openrouter, "chat_completion", lambda *a, **k: json.dumps([{"instruction": "Q", "output": "A"}])
    )
    recipe_id = seed_recipe(service, "Passage.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    updated = service.update_record(
        recipe_id, 0, RecipeRecordUpdate(record={"instruction": "Edited", "output": "New"})
    )
    assert updated.record["instruction"] == "Edited"
    assert updated.generator == "llm"  # provenance preserved through edit


def test_add_and_delete_records(settings: Settings):
    service, _ = make_service(settings)
    recipe_id = seed_recipe(service, "Passage.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="rules"))
    added = service.add_record(
        recipe_id, RecipeRecordCreate(record={"instruction": "Manual", "output": "Answer"})
    )
    assert added.record["instruction"] == "Manual"
    recipe = service.delete_records(recipe_id, [added.index])
    page = service.list_records(recipe_id, page=1, page_size=50)
    assert all(record.record.get("instruction") != "Manual" for record in page.records)
    assert recipe.status == "ready"


# ---- commit -----------------------------------------------------------------


def test_commit_creates_llm_dataset(settings: Settings):
    service, datasets = make_service(settings)
    recipe_id = seed_recipe(service, "What is AI?\nArtificial intelligence.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="rules"))
    response = service.commit(recipe_id, RecipeCommitRequest(name="From Docs"))
    assert response.committed_records > 0
    dataset = response.dataset
    assert dataset.task_type == "llm_finetune"
    assert dataset.origin == "recipe"
    assert dataset.origin_ref == recipe_id
    # Records land in the unassigned inbox.
    assert dataset.splits["unassigned"].item_count == response.committed_records


def test_commit_with_zero_records_blocked(settings: Settings):
    service, _ = make_service(settings)
    recipe = service.create_recipe(RecipeCreate(name="Empty"))
    with pytest.raises(Exception) as exc:
        service.commit(recipe.id, RecipeCommitRequest())
    assert "at least one record" in str(exc.value)


def test_generate_with_no_sources_rejected(settings: Settings):
    service, _ = make_service(settings)
    recipe = service.create_recipe(RecipeCreate(name="No sources"))
    with pytest.raises(Exception) as exc:
        service.generate(recipe.id, RecipeGenerateRequest(mode="rules"))
    assert "No extractable text" in str(exc.value)


# ---- OpenRouter key masking -------------------------------------------------


def test_openrouter_key_masked_in_settings(make_settings, tmp_path: Path):
    env_path = tmp_path / ".env"
    settings = make_settings()
    service = SettingsService(settings, env_path=env_path)
    service.save_openrouter_key("sk-secret-value")
    assert "OPENROUTER_API_KEY=sk-secret-value" in env_path.read_text()
    # The read model exposes only a boolean, never the raw key.
    from app.schemas import PlatformSettingsRead

    assert "openrouter_api_key" not in PlatformSettingsRead.model_fields
    assert "openrouter_api_key_configured" in PlatformSettingsRead.model_fields


def test_key_never_written_to_recipe_manifest(make_settings, monkeypatch):
    settings = _llm_settings(make_settings)
    service, _ = make_service(settings)
    monkeypatch.setattr(
        openrouter, "chat_completion", lambda *a, **k: json.dumps([{"instruction": "Q", "output": "A"}])
    )
    recipe_id = seed_recipe(service, "Passage about a topic.")
    service.generate(recipe_id, RecipeGenerateRequest(mode="llm"))
    root = service.storage.recipes / recipe_id
    manifest_text = (root / "manifest.json").read_text()
    records_text = (root / "records.jsonl").read_text()
    assert "sk-test-key" not in manifest_text
    assert "sk-test-key" not in records_text


def test_service_class_alias():
    # `ServiceCls` and the package export are the same class.
    assert ServiceCls is RecipeService
